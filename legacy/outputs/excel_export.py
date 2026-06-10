"""
Excel export — writes the TRS deal pipeline to an .xlsx workbook.

On every run it rewrites the workbook from SQLite with four sheets:
  * Tier 1 — Priority   (gold highlight, sorted by score desc)
  * Tier 2 — Watch List (slate highlight)
  * All Records         (every company)
  * Run Log             (timestamped run history)

Header rows are frozen and auto-filtered. Score / OSHA severity / tier columns
are color-coded; listing URLs are clickable. If the target file is locked
(open in Excel on Windows), it writes a timestamped temp copy and logs a
warning instead of crashing. Output path comes from config.EXCEL_OUTPUT_PATH.
"""

import logging
import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import EXCEL_OUTPUT_PATH
from database.db import Database

# TRS brand colors
NAVY = "1B365D"
GOLD = "D4A843"
SLATE = "94A3B8"
WHITE = "FFFFFF"
LIGHT_GOLD = "FDF3D0"
LIGHT_SLATE = "F1F4F7"

# (header label, db field, column width, number format)
COLUMNS = [
    ("Company Name", "company_name", 30, None),
    ("Tier", "tier", 8, None),
    ("Score", "score_total", 7, None),
    ("Stage", "pipeline_stage", 14, None),
    ("Source", "source", 12, None),
    ("State", "state", 6, None),
    ("City", "city", 16, None),
    ("Sector", "trs_sector", 16, None),
    ("Subsector", "trs_subsector", 20, None),
    ("SIC", "sic_code", 7, None),
    ("NAICS", "naics_code", 8, None),
    ("Rev Est Low", "revenue_estimate_low", 14, '"$"#,##0'),
    ("Rev Est High", "revenue_estimate_high", 14, '"$"#,##0'),
    ("EBITDA Est", "ebitda_estimate", 14, '"$"#,##0'),
    ("Employees (High)", "employee_count_high", 14, None),
    ("Founded", "founded_year", 9, None),
    ("Filing Status", "filing_status", 12, None),
    ("OSHA Severity", "osha_severity", 13, None),
    ("OSHA Violations", "osha_violation_count", 14, None),
    ("Days on Market", "days_on_market", 14, None),
    ("Asking Price", "asking_price", 14, '"$"#,##0'),
    ("Listed Revenue", "listed_revenue", 14, '"$"#,##0'),
    ("Listed Cash Flow", "listed_cash_flow", 14, '"$"#,##0'),
    ("Broker", "broker_name", 18, None),
    ("Listing URL", "listing_url", 30, None),
    ("Assigned To", "assigned_to", 12, None),
    ("Next Action", "next_action", 24, None),
    ("Next Action Date", "next_action_date", 14, None),
    ("Date Added", "date_added", 12, None),
    ("Notes", "notes", 40, None),
]

_OSHA_COLORS = {"high": "B71C1C", "medium": "E65100", "low": "F9A825", "none": "2E7D32"}
_TIER_LABELS = {1: "Tier 1", 2: "Tier 2", 3: "Tier 3"}


class ExcelExport:
    def __init__(self, db=None):
        self.db = db or Database()

    def write(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        tier1 = self.db.get_by_tier(1)
        tier2 = self.db.get_by_tier(2)
        self._build_sheet(wb, "Tier 1 — Priority", tier1)
        self._build_sheet(wb, "Tier 2 — Watch List", tier2)
        self._build_sheet(wb, "All Records", self.db.get_all())
        self._build_run_log_sheet(wb, self.db.get_run_log())

        output_path = self._resolve_path()
        wb.save(output_path)

        exported_ids = [r["id"] for r in (tier1 + tier2) if r.get("id")]
        self.db.set_excel_exported(exported_ids)

        logging.info(
            "Excel export written: %s (%d T1, %d T2)",
            output_path,
            len(tier1),
            len(tier2),
        )
        print(f"  Excel written: {output_path}")
        return output_path

    def _resolve_path(self):
        path = str(EXCEL_OUTPUT_PATH)
        if os.path.exists(path):
            try:
                with open(path, "r+b"):
                    pass
            except PermissionError:
                stamp = datetime.now().strftime("%H%M%S")
                path = path.replace(".xlsx", f"_temp_{stamp}.xlsx")
                logging.warning("Excel file locked; writing to %s instead.", path)
        return path

    def _build_sheet(self, wb, sheet_name, rows):
        ws = wb.create_sheet(title=sheet_name)
        header_fill = PatternFill("solid", fgColor=NAVY)
        header_font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        header_align = Alignment(horizontal="center", vertical="center")

        for col_idx, (label, _field, width, _fmt) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.row_dimensions[1].height = 20
        ws.freeze_panes = "A2"

        for row_idx, record in enumerate(rows, start=2):
            tier = record.get("tier", 3)
            row_fill = None
            if tier == 1:
                row_fill = PatternFill("solid", fgColor=LIGHT_GOLD)
            elif tier == 2:
                row_fill = PatternFill("solid", fgColor=LIGHT_SLATE)

            for col_idx, (_label, field, _width, num_fmt) in enumerate(COLUMNS, 1):
                value = record.get(field)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment(vertical="top")
                cell.border = Border(bottom=Side(style="hair", color="E0E0E0"))
                if row_fill:
                    cell.fill = row_fill
                if num_fmt and value is not None:
                    cell.number_format = num_fmt

                if field == "listing_url" and value:
                    cell.hyperlink = value
                    cell.font = Font(
                        name="Calibri", size=10, color="0563C1", underline="single"
                    )
                elif field == "score_total" and value is not None:
                    if value >= 70:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="1B5E20")
                    elif value >= 45:
                        cell.font = Font(name="Calibri", size=10, color="E65100")
                elif field == "osha_severity" and value:
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=(value == "high"),
                        color=_OSHA_COLORS.get(value, "000000"),
                    )
                elif field == "tier" and value is not None:
                    cell.value = _TIER_LABELS.get(value, str(value))
                    if value == 1:
                        cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
                    elif value == 2:
                        cell.font = Font(name="Calibri", size=10, color="555555")

        ws.auto_filter.ref = ws.dimensions
        tab_colors = {"Tier 1 — Priority": GOLD, "Tier 2 — Watch List": SLATE}
        if sheet_name in tab_colors:
            ws.sheet_properties.tabColor = tab_colors[sheet_name]

    def _build_run_log_sheet(self, wb, run_log_rows):
        ws = wb.create_sheet(title="Run Log")
        ws.sheet_properties.tabColor = "CCCCCC"
        headers = [
            "Run Date", "Source", "Records Pulled", "New Records", "Updated",
            "Tier 1", "Tier 2", "Errors", "Duration (s)",
        ]
        fields = [
            "run_date", "source", "records_pulled", "records_new",
            "records_updated", "tier1_count", "tier2_count", "errors",
            "duration_seconds",
        ]
        header_fill = PatternFill("solid", fgColor=NAVY)
        header_font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        for col_idx, label in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill = header_fill
            cell.font = header_font
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

        for row_idx, log in enumerate(run_log_rows, start=2):
            for col_idx, field in enumerate(fields, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=log.get(field))
                cell.font = Font(name="Calibri", size=10)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
