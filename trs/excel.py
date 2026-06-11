"""Excel workbook export (openpyxl).

Tabs: Summary dashboard, Scored Targets, New This Run, Watch List (60-79), Source
Log. House style: navy (#1B365D) header fills, Georgia bold headings, Calibri body,
blue font for any manually editable cell. No em dashes anywhere in cell text.
"""

from __future__ import annotations

import json
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db, settings

NAVY = "1B365D"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Georgia", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Georgia", bold=True, color=NAVY, size=16)
SUBTITLE_FONT = Font(name="Georgia", bold=True, color=NAVY, size=12)
BODY_FONT = Font(name="Calibri", size=10)
EDITABLE_FONT = Font(name="Calibri", size=10, color="0000CC")  # blue = user-editable
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def _no_emdash(value):
    if isinstance(value, str):
        return value.replace("—", "-").replace("–", "-")
    return value


def _header_row(ws, headers, row=1):
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=_no_emdash(text))
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


# Columns shared by the three target tabs. (header, db/derived key, editable?)
TARGET_COLUMNS = [
    ("Company", "name", False),
    ("Score", "total", False),
    ("Niche /30", "niche_fit", False),
    ("Size /25", "size_fit", False),
    ("Succession /25", "succession", False),
    ("Geo /10", "geography", False),
    ("Confidence /10", "data_confidence", False),
    ("State", "state", False),
    ("NAICS", "naics_code", False),
    ("Employees", "employee_count", False),
    ("Revenue est.", "revenue_estimate", False),
    ("Founded", "founded_year", False),
    ("New?", "new_flag", False),
    ("Sources", "sources", False),
    ("Pipeline Stage", "_edit_stage", True),
    ("Owner / Contact (verify)", "_edit_owner", True),
    ("Notes", "_edit_notes", True),
]


def _gather_rows(conn):
    run_id = db.latest_run_id(conn)
    rows = []
    for comp in db.all_companies(conn, include_excluded=False):
        score = conn.execute(
            "SELECT * FROM scores WHERE company_id=? ORDER BY run_id DESC LIMIT 1",
            (comp["id"],),
        ).fetchone()
        sigs = db.signals_for(conn, comp["id"])
        urls = sorted({s["source_url"] for s in sigs if s["source_url"]})
        rev = comp["revenue_estimate"]
        row = {
            "name": comp["name"],
            "state": comp["state"] or "",
            "naics_code": comp["naics_code"] or "",
            "employee_count": comp["employee_count"] or "",
            "revenue_estimate": f"${rev:,.0f}" if rev else "",
            "founded_year": comp["founded_year"] or "",
            "new_flag": "Yes" if comp["new_this_run"] else "",
            "sources": " ; ".join(urls[:3]),
            "total": score["total"] if score else 0,
            "niche_fit": score["niche_fit"] if score else 0,
            "size_fit": score["size_fit"] if score else 0,
            "succession": score["succession"] if score else 0,
            "geography": score["geography"] if score else 0,
            "data_confidence": score["data_confidence"] if score else 0,
            "_new": bool(comp["new_this_run"]),
        }
        rows.append(row)
    rows.sort(key=lambda r: r["total"], reverse=True)
    return rows


def _write_target_sheet(ws, rows):
    _header_row(ws, [h for h, _, _ in TARGET_COLUMNS])
    for i, row in enumerate(rows, start=2):
        for col, (_, key, editable) in enumerate(TARGET_COLUMNS, start=1):
            val = "" if key.startswith("_edit") else row.get(key, "")
            c = ws.cell(row=i, column=col, value=_no_emdash(val))
            c.font = EDITABLE_FONT if editable else BODY_FONT
            c.alignment = WRAP
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TARGET_COLUMNS))}{max(len(rows) + 1, 1)}"
    _autosize(ws, {1: 34, 2: 8, 3: 9, 4: 9, 5: 13, 6: 8, 7: 13, 8: 7,
                   9: 9, 10: 11, 11: 14, 12: 9, 13: 7, 14: 46, 15: 16,
                   16: 26, 17: 30})


def _write_summary(ws, conn, rows, universe):
    ws["A1"] = _no_emdash("TRS Investments - Deal Sourcing Summary")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Run date: {date.today().isoformat()}"
    ws["A2"].font = BODY_FONT

    scored = [r for r in rows if r["total"] > 0]
    w = settings.weights()
    surface = w.get("surface_threshold", 50)
    wl_low, wl_high = w["watch_list"]["low"], w["watch_list"]["high"]

    metrics = [
        ("Companies in database (in scope)", len(rows)),
        ("Scored targets (>= surface threshold)", len([r for r in scored if r["total"] >= surface])),
        (f"Watch List ({wl_low}-{wl_high})", len([r for r in scored if wl_low <= r["total"] <= wl_high])),
        ("New this run", len([r for r in rows if r["_new"]])),
        ("Top score", max((r["total"] for r in scored), default=0)),
    ]
    ws["A4"] = "Pipeline metrics"
    ws["A4"].font = SUBTITLE_FONT
    r = 5
    for label, value in metrics:
        ws.cell(row=r, column=1, value=_no_emdash(label)).font = BODY_FONT
        ws.cell(row=r, column=2, value=value).font = Font(name="Calibri", bold=True, size=10)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Census universe map (CBP)").font = SUBTITLE_FONT
    r += 1
    if universe:
        for label, value in [
            ("NAICS / states", _no_emdash(universe.get("scope", ""))),
            ("Establishments (county cells)", universe.get("total_establishments", 0)),
            ("Total employees", universe.get("total_employees", 0)),
            ("Counties with data", universe.get("counties", 0)),
            ("Size-suppressed cells", universe.get("suppressed_cells", 0)),
        ]:
            ws.cell(row=r, column=1, value=_no_emdash(label)).font = BODY_FONT
            ws.cell(row=r, column=2, value=value).font = BODY_FONT
            r += 1
    else:
        ws.cell(row=r, column=1, value="Not run this pass.").font = BODY_FONT

    _autosize(ws, {1: 40, 2: 30})


def _write_source_log(ws, conn, universe):
    _header_row(ws, ["Run date", "Source", "Records pulled", "New", "Notes"])
    logs = conn.execute(
        "SELECT * FROM run_log ORDER BY id DESC LIMIT 50"
    ).fetchall()
    for i, lg in enumerate(logs, start=2):
        for col, key in enumerate(
                ["run_date", "source", "records_pulled", "records_new", "notes"], start=1):
            c = ws.cell(row=i, column=col, value=_no_emdash(lg[key]))
            c.font = BODY_FONT
            c.alignment = WRAP
    _autosize(ws, {1: 22, 2: 12, 3: 14, 4: 8, 5: 80})


def export(conn, *, universe: dict | None = None, path=None) -> str:
    path = str(path or (settings.OUTPUTS_DIR /
              f"TRS_Deal_Pipeline_{date.today().isoformat()}.xlsx"))
    rows = _gather_rows(conn)
    w = settings.weights()
    surface = w.get("surface_threshold", 50)
    wl_low, wl_high = w["watch_list"]["low"], w["watch_list"]["high"]

    wb = Workbook()
    _write_summary(wb.active, conn, rows, universe)
    wb.active.title = "Summary"

    _write_target_sheet(wb.create_sheet("Scored Targets"),
                         [r for r in rows if r["total"] >= surface])
    _write_target_sheet(wb.create_sheet("New This Run"),
                         [r for r in rows if r["_new"]])
    _write_target_sheet(wb.create_sheet("Watch List"),
                         [r for r in rows if wl_low <= r["total"] <= wl_high])
    _write_source_log(wb.create_sheet("Source Log"), conn, universe)

    wb.save(path)
    print(f"[excel] wrote {path}")
    return path
